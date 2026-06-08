#!/usr/bin/env python3
"""
Signing Agent HQ — free lead magnet generator.

The top-of-funnel giveaway promised by the landing page, email sequence and
video script: a free IRS mileage log for notaries. Builds both versions
described in validation/.../lead-magnet-mileage-log.md:

  dist/lead-magnet/Notary-Mileage-Log.xlsx            (Google Sheets / Excel, live formulas)
  dist/lead-magnet/Notary-Mileage-Log-printable.pdf   (blank, for the glovebox)

Both carry a soft upsell bridge to the paid Signing Agent HQ.

Usage:  python3 build/build_lead_magnet.py
"""

from __future__ import annotations

import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# brand palette (matches build_cover.py / build_workbook.py)
NAVY = "1F2A44"
TEAL = "1F7A8C"
GOLD = "E0A458"
WHITE = "FFFFFF"
INK = "1A2238"

OUT_DIR = "dist/lead-magnet"
RATE_DEFAULT = 0.70  # 2025 IRS standard rate; user updates each year
ROWS = 40

CURRENCY = '"$"#,##0.00'
RATE_FMT = '"$"#,##0.000'
MILES = "#,##0.0"

thin = Side(style="thin", color="C7D0DA")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def fill(hexcolor):
    return PatternFill("solid", fgColor=hexcolor)


HEADERS = ["Date", "Order # / Purpose", "From", "To", "Round-trip miles", "Deduction ($)"]
WIDTHS = [13, 26, 16, 16, 16, 16]


def build_xlsx():
    wb = Workbook()
    ws = wb.active
    ws.title = "Mileage Log"
    ws.sheet_view.showGridLines = False
    for c, w in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w

    last = get_column_letter(len(HEADERS))
    # title
    ws.merge_cells(f"A1:{last}1")
    ws["A1"] = "The Notary Mileage Log"
    ws["A1"].font = Font(name="Calibri", size=22, bold=True, color=WHITE)
    ws["A1"].fill = fill(NAVY)
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 40
    ws.merge_cells(f"A2:{last}2")
    ws["A2"] = "Never miss a deductible mile again — track every trip in 10 seconds."
    ws["A2"].font = Font(name="Calibri", size=12, bold=True, color=WHITE)
    ws["A2"].fill = fill(TEAL)
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[2].height = 22

    # rate parameter
    ws.merge_cells("A4:D4")
    ws["A4"] = "IRS standard mileage rate (this year, $/mile):"
    ws["A4"].font = Font(name="Calibri", size=11, bold=True, color=NAVY)
    ws["A4"].alignment = Alignment(horizontal="right", vertical="center")
    rate = ws["E4"]
    rate.value = RATE_DEFAULT
    rate.number_format = RATE_FMT
    rate.fill = fill(GOLD)
    rate.font = Font(bold=True, color=NAVY)
    rate.alignment = CENTER
    rate.border = BORDER
    ws["F4"] = "← enter the official rate"
    ws["F4"].font = Font(name="Calibri", size=9, italic=True, color="566173")
    ws["F4"].alignment = Alignment(horizontal="left", vertical="center")

    # header row
    hr = 6
    for c, text in enumerate(HEADERS, start=1):
        cell = ws.cell(row=hr, column=c, value=text)
        cell.fill = fill(TEAL)
        cell.font = Font(bold=True, color=WHITE)
        cell.alignment = CENTER
        cell.border = BORDER
    ws.freeze_panes = f"A{hr + 1}"

    first = hr + 1
    last_row = hr + ROWS
    for r in range(first, last_row + 1):
        for c in range(1, len(HEADERS) + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            cell.alignment = CENTER
            if (r - first) % 2:
                cell.fill = fill("F4F6FA")
        ws.cell(row=r, column=5).number_format = MILES
        ded = ws.cell(row=r, column=6)
        ded.value = f"=IF(E{r}=\"\",\"\",E{r}*$E$4)"
        ded.number_format = CURRENCY

    # totals
    tr = last_row + 1
    ws.merge_cells(f"A{tr}:D{tr}")
    ws.cell(row=tr, column=1, value="TOTAL  (year to date)")
    ws.cell(row=tr, column=1).font = Font(bold=True, color=NAVY)
    ws.cell(row=tr, column=1).alignment = Alignment(horizontal="right", vertical="center")
    tm = ws.cell(row=tr, column=5, value=f"=SUM(E{first}:E{last_row})")
    tm.number_format = MILES
    td = ws.cell(row=tr, column=6, value=f"=SUM(F{first}:F{last_row})")
    td.number_format = CURRENCY
    for c in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=tr, column=c)
        cell.fill = fill("E8EEF5")
        cell.border = BORDER
        if c in (5, 6):
            cell.font = Font(bold=True, color=TEAL, size=12)

    # disclaimer
    note_r = tr + 2
    ws.merge_cells(f"A{note_r}:{last}{note_r}")
    ws.cell(row=note_r, column=1,
            value="Keep this log contemporaneous — the IRS requires a timely, accurate record. "
                  "This is an organizer, not tax advice.")
    ws.cell(row=note_r, column=1).font = Font(size=9.5, italic=True, color="566173")
    ws.cell(row=note_r, column=1).alignment = LEFT

    # upsell bridge
    up_r = note_r + 2
    ws.merge_cells(f"A{up_r}:{last}{up_r}")
    ws.cell(row=up_r, column=1,
            value="Love this? Signing Agent HQ turns it into your whole business: "
                  "signing CRM, net profit per job, invoices, and a one-click tax summary. "
                  "→ Get it at founder price $19.")
    ws.cell(row=up_r, column=1).font = Font(size=11, bold=True, color=NAVY)
    ws.cell(row=up_r, column=1).fill = fill(GOLD)
    ws.cell(row=up_r, column=1).alignment = LEFT
    ws.row_dimensions[up_r].height = 38

    os.makedirs(OUT_DIR, exist_ok=True)
    out = os.path.join(OUT_DIR, "Notary-Mileage-Log.xlsx")
    wb.save(out)
    print("Wrote", out)


def build_pdf():
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import (Paragraph, SimpleDocTemplate, Spacer, Table,
                                    TableStyle)

    NAVYc = colors.HexColor("#1F2A44")
    TEALc = colors.HexColor("#1F7A8C")
    GOLDc = colors.HexColor("#E0A458")
    GREY = colors.HexColor("#566173")

    styles = getSampleStyleSheet()
    title = ParagraphStyle("t", parent=styles["Title"], fontSize=24, textColor=NAVYc,
                           spaceAfter=2, leading=27, alignment=0)
    sub = ParagraphStyle("s", parent=styles["Normal"], fontSize=12, textColor=TEALc,
                         spaceAfter=14)
    rate = ParagraphStyle("r", parent=styles["Normal"], fontSize=12, textColor=NAVYc,
                          spaceAfter=12)
    note = ParagraphStyle("n", parent=styles["Normal"], fontSize=8.5, textColor=GREY,
                          spaceBefore=10, leading=12)
    up = ParagraphStyle("u", parent=styles["Normal"], fontSize=10.5, textColor=NAVYc,
                        leading=14, spaceBefore=8)

    story = [
        Paragraph("The Notary Mileage Log", title),
        Paragraph("Never miss a deductible mile again.", sub),
        Paragraph("IRS standard mileage rate (this year): "
                  "<b>$ ____________ / mile</b>", rate),
    ]

    headers = ["Date", "Order # / Purpose", "From", "To", "Miles", "Deduction $"]
    col_w = [0.95, 2.25, 1.15, 1.15, 0.7, 1.0]
    col_w = [w * inch for w in col_w]
    data = [headers] + [[""] * 6 for _ in range(22)]
    data.append(["TOTAL (YTD)", "", "", "", "", ""])

    t = Table(data, colWidths=col_w, rowHeights=[0.32 * inch] + [0.3 * inch] * 22 + [0.34 * inch])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), TEALc),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 9.5),
        ("ALIGN", (0, 0), (-1, 0), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.6, colors.HexColor("#C7D0DA")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#F4F6FA")]),
        ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E8EEF5")),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("SPAN", (0, -1), (3, -1)),
        ("ALIGN", (0, -1), (0, -1), "RIGHT"),
    ]))
    story.append(t)

    story += [
        Paragraph("Keep this log contemporaneous — the IRS requires a timely, accurate "
                  "record. This is an organizer, not tax advice.", note),
        Paragraph("<b>Love this?</b> Signing Agent HQ turns it into your whole business: "
                  "signing CRM, net profit per job, invoices, and a one-click Schedule&nbsp;C "
                  "tax summary. &rarr; Get it at founder price <b>$19</b>.", up),
    ]

    os.makedirs(OUT_DIR, exist_ok=True)
    out = os.path.join(OUT_DIR, "Notary-Mileage-Log-printable.pdf")
    doc = SimpleDocTemplate(out, pagesize=LETTER, topMargin=0.7 * inch,
                            bottomMargin=0.6 * inch, leftMargin=0.7 * inch,
                            rightMargin=0.7 * inch, title="The Notary Mileage Log")
    doc.build(story)
    print("Wrote", out)


if __name__ == "__main__":
    build_xlsx()
    build_pdf()
