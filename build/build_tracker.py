#!/usr/bin/env python3
"""
Signing Agent HQ — validation tracker generator.

A simple operator spreadsheet to log the validation metrics against the
Go/No-Go thresholds from validation-plan.md, with an automatic verdict. The
one tool the seller touches daily during the 14-day launch window.

Output:  dist/Validation-Tracker.xlsx
Usage:   python3 build/build_tracker.py
"""

from __future__ import annotations

import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

NAVY = "1F2A44"
TEAL = "1F7A8C"
GOLD = "E0A458"
GREEN = "2E7D32"
WHITE = "FFFFFF"

H1 = Font(name="Calibri", size=20, bold=True, color=WHITE)
H2 = Font(name="Calibri", size=12, bold=True, color=WHITE)
HDR = Font(name="Calibri", size=11, bold=True, color=WHITE)
BOLD = Font(name="Calibri", size=11, bold=True, color=NAVY)
NOTE = Font(name="Calibri", size=10, italic=True, color="566173")
BIG = Font(name="Calibri", size=18, bold=True, color=NAVY)

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
thin = Side(style="thin", color="C7D0DA")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def fill(c):
    return PatternFill("solid", fgColor=c)


def banner(ws, title, sub, span):
    last = get_column_letter(span)
    ws.merge_cells(f"A1:{last}1")
    ws["A1"] = title
    ws["A1"].fill = fill(NAVY)
    ws["A1"].font = H1
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 38
    ws.merge_cells(f"A2:{last}2")
    ws["A2"] = sub
    ws["A2"].fill = fill(TEAL)
    ws["A2"].font = H2
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[2].height = 22


def hrow(ws, row, headers):
    for c, t in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=c, value=t)
        cell.fill = fill(TEAL)
        cell.font = HDR
        cell.alignment = CENTER
        cell.border = BORDER


def build():
    wb = Workbook()
    ws = wb.active
    ws.title = "Validation Tracker"
    ws.sheet_view.showGridLines = False
    banner(ws, "Validation Tracker", "Log your numbers. The verdict updates itself.", 5)

    # --- scoreboard ---
    ws["A4"] = "Scoreboard (enter your numbers in the blue column)"
    ws["A4"].font = BOLD
    hrow(ws, 5, ["Metric", "GO threshold", "Your number", "Status", "Notes"])
    # metric, threshold value, is_percent
    metrics = [
        ("Landing clicks", 50, False),
        ("Lead-magnet opt-in rate", 0.20, True),
        ("Email / waitlist", 20, False),
        ("Pre-orders @ $19", 5, False),
        ("Qualitative replies", 8, False),
    ]
    first = 6
    for i, (name, thr, pct) in enumerate(metrics):
        r = first + i
        ws.cell(row=r, column=1, value=name).font = Font(bold=(name == "Pre-orders @ $19"), color=NAVY)
        tcell = ws.cell(row=r, column=2, value=thr)
        ucell = ws.cell(row=r, column=3)  # user input (blank)
        ucell.fill = fill("EAF2FB")
        scell = ws.cell(row=r, column=4,
                        value=f'=IF(C{r}="","—",IF(C{r}>=B{r},"GO","not yet"))')
        if pct:
            tcell.number_format = "0%"
            ucell.number_format = "0%"
        scell.alignment = CENTER
        scell.font = Font(bold=True, color=NAVY)
        for c in range(1, 6):
            ws.cell(row=r, column=c).border = BORDER

    # --- verdict ---
    pre_row = first + 3  # Pre-orders row
    vr = first + len(metrics) + 1
    ws.merge_cells(f"A{vr}:B{vr}")
    ws.cell(row=vr, column=1, value="DECISION (gated on pre-orders)").font = BOLD
    verdict = ws.cell(row=vr, column=3,
                      value=f'=IF(C{pre_row}="","Awaiting data",'
                            f'IF(C{pre_row}>=B{pre_row},"GO — build & deliver",'
                            f'IF(C{first+2}>=B{first+2},"REPOSITION (narrow the wedge)","KEEP PUSHING")))')
    ws.merge_cells(f"C{vr}:E{vr}")
    verdict.font = BIG
    verdict.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    verdict.fill = fill(GOLD)
    ws.row_dimensions[vr].height = 30
    for c in range(1, 6):
        ws.cell(row=vr, column=c).border = BORDER

    ws.cell(row=vr + 2, column=1,
            value="Rule: Pre-orders ≥ 5 → GO. Waitlist ≥ 20 but pre-orders < 5 → reposition "
                  "to a tighter wedge before discarding. (See validation-plan.md.)").font = NOTE
    ws.merge_cells(f"A{vr+2}:E{vr+2}")

    # --- daily log ---
    lr = vr + 4
    ws.cell(row=lr, column=1, value="Daily log").font = BOLD
    hrow(ws, lr + 1, ["Date", "Channel / action", "Clicks", "Opt-ins", "Pre-orders"])
    for r in range(lr + 2, lr + 2 + 21):
        for c in range(1, 6):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            if (r - (lr + 2)) % 2:
                cell.fill = fill("F4F6FA")
        ws.cell(row=r, column=1).number_format = "yyyy-mm-dd"
    # totals
    tr = lr + 2 + 21
    ws.cell(row=tr, column=2, value="Totals").font = BOLD
    for c, col in [(3, "C"), (4, "D"), (5, "E")]:
        tc = ws.cell(row=tr, column=c, value=f"=SUM({col}{lr+2}:{col}{tr-1})")
        tc.font = Font(bold=True, color=TEAL)
        tc.fill = fill("E8EEF5")
        tc.border = BORDER
    ws.cell(row=tr, column=2).fill = fill("E8EEF5")
    ws.cell(row=tr, column=2).border = BORDER

    # widths
    for col, w in {"A": 16, "B": 16, "C": 16, "D": 14, "E": 22}.items():
        ws.column_dimensions[col].width = w

    os.makedirs("dist", exist_ok=True)
    out = "dist/Validation-Tracker.xlsx"
    wb.save(out)
    print("Wrote", out)


if __name__ == "__main__":
    build()
